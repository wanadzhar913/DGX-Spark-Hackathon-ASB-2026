"""
Generate a hackathon pitch deck (PowerPoint) for the Live AI Video Monitoring
system. Dark theme matching the dashboard. Embeds the dashboard screenshot if
one is present in docs/screenshots/.

Usage (from repo root):
    backend/.venv/Scripts/python scripts/make_slides.py
Output:
    docs/DGX-Spark-AI-Video-Monitoring.pptx
"""

import os

from pptx import Presentation
from pptx.dml.color import RGBColor
from pptx.enum.text import MSO_ANCHOR, PP_ALIGN
from pptx.util import Emu, Inches, Pt

HERE = os.path.dirname(os.path.abspath(__file__))
REPO = os.path.abspath(os.path.join(HERE, ".."))
SHOTS = os.path.join(REPO, "docs", "screenshots")
OUT = os.path.join(REPO, "docs", "DGX-Spark-AI-Video-Monitoring.pptx")

# ---- palette (matches the dashboard) -------------------------------------
BG = RGBColor(0x0B, 0x0F, 0x14)
PANEL = RGBColor(0x12, 0x18, 0x20)
CARD = RGBColor(0x1A, 0x22, 0x2D)
TEXT = RGBColor(0xE6, 0xED, 0xF3)
MUTED = RGBColor(0x8A, 0x97, 0xA6)
ACCENT = RGBColor(0x38, 0xBD, 0xF8)
GREEN = RGBColor(0x76, 0xB9, 0x00)  # NVIDIA-ish green
RED = RGBColor(0xEF, 0x44, 0x44)
AMBER = RGBColor(0xF5, 0x9E, 0x0B)

# 16:9
EMU_W = Inches(13.333)
EMU_H = Inches(7.5)

prs = Presentation()
prs.slide_width = EMU_W
prs.slide_height = EMU_H
BLANK = prs.slide_layouts[6]


def slide():
    s = prs.slides.add_slide(BLANK)
    bg = s.shapes.add_shape(1, 0, 0, EMU_W, EMU_H)  # rectangle
    bg.fill.solid()
    bg.fill.fore_color.rgb = BG
    bg.line.fill.background()
    bg.shadow.inherit = False
    # send to back
    sp = bg._element
    sp.getparent().remove(sp)
    s.shapes._spTree.insert(2, sp)
    return s


def box(s, x, y, w, h, fill=None, line=None, line_w=1.0, radius=False):
    shp = s.shapes.add_shape(5 if radius else 1, Inches(x), Inches(y),
                             Inches(w), Inches(h))
    if fill is None:
        shp.fill.background()
    else:
        shp.fill.solid()
        shp.fill.fore_color.rgb = fill
    if line is None:
        shp.line.fill.background()
    else:
        shp.line.color.rgb = line
        shp.line.width = Pt(line_w)
    shp.shadow.inherit = False
    return shp


def text(s, x, y, w, h, runs, align=PP_ALIGN.LEFT, anchor=MSO_ANCHOR.TOP,
         space_after=6):
    """runs: list of (string, size, color, bold)."""
    tb = s.shapes.add_textbox(Inches(x), Inches(y), Inches(w), Inches(h))
    tf = tb.text_frame
    tf.word_wrap = True
    tf.vertical_anchor = anchor
    for i, line in enumerate(runs):
        p = tf.paragraphs[0] if i == 0 else tf.add_paragraph()
        p.alignment = align
        p.space_after = Pt(space_after)
        if isinstance(line, tuple):
            line = [line]
        for (txt, size, color, bold) in line:
            r = p.add_run()
            r.text = txt
            r.font.size = Pt(size)
            r.font.color.rgb = color
            r.font.bold = bold
            r.font.name = "Segoe UI"
    return tb


def bullets(s, x, y, w, h, items, size=16, color=TEXT, gap=10):
    tb = s.shapes.add_textbox(Inches(x), Inches(y), Inches(w), Inches(h))
    tf = tb.text_frame
    tf.word_wrap = True
    for i, it in enumerate(items):
        p = tf.paragraphs[0] if i == 0 else tf.add_paragraph()
        p.space_after = Pt(gap)
        r = p.add_run()
        r.text = "•  " + it
        r.font.size = Pt(size)
        r.font.color.rgb = color
        r.font.name = "Segoe UI"
    return tb


def kicker(s, label):
    bar = box(s, 0.7, 0.62, 0.12, 0.5, fill=ACCENT)
    text(s, 0.95, 0.5, 11, 0.6, [[(label, 14, ACCENT, True)]])


# ==========================================================================
# Slide 1 — Title
# ==========================================================================
s = slide()
box(s, 0, 0, 13.333, 0.18, fill=GREEN)
text(s, 0.9, 2.2, 11.5, 1.6, [
    [("Live AI Video Monitoring", 46, TEXT, True)],
    [("Real-time behavior detection across many camera feeds — on the DGX Spark",
      20, MUTED, False)],
])
# chips
chips = ["Multi-feed CCTV wall", "AI behavior events", "Runs on DGX Spark"]
cx = 0.9
for c in chips:
    w = 0.26 + 0.115 * len(c)
    box(s, cx, 4.2, w, 0.5, fill=CARD, line=ACCENT, line_w=1.0, radius=True)
    text(s, cx, 4.26, w, 0.4, [[(c, 13, ACCENT, True)]], align=PP_ALIGN.CENTER)
    cx += w + 0.25
text(s, 0.9, 6.7, 11.5, 0.5,
     [[("DGX Spark Hackathon · ASB 2026", 13, MUTED, False)]])

# ==========================================================================
# Slide 2 — Problem
# ==========================================================================
s = slide()
kicker(s, "THE PROBLEM")
text(s, 0.9, 1.1, 11.5, 1.0, [[("Watching feeds doesn't scale", 34, TEXT, True)]])
bullets(s, 0.95, 2.3, 11.4, 4, [
    "Security teams stare at walls of camera feeds — attention fades, incidents are missed.",
    "Anti-social behavior (loitering, vandalism, fights, intrusion) is easy to miss in real time.",
    "Cloud video AI means streaming sensitive footage off-site: cost, latency, privacy risk.",
    "Operators need the crucial moments surfaced, not hours of raw video.",
], size=19, gap=16)

# ==========================================================================
# Slide 3 — Solution
# ==========================================================================
s = slide()
kicker(s, "OUR SOLUTION")
text(s, 0.9, 1.1, 11.5, 1.0, [[("An AI monitoring wall that flags what matters",
                                30, TEXT, True)]])
cards = [
    ("Multi-feed wall", "Up to 9 live camera feeds in one grid, looping or live.", ACCENT),
    ("AI behavior events", "A local vision model watches a configurable watchlist and raises events.", GREEN),
    ("Crucial info panel", "Each event: camera, time, type, confidence, severity + plain-language description.", AMBER),
    ("Snapshot evidence", "A thumbnail of the exact frame where each event happened.", RED),
]
x = 0.9
for title, body, col in cards:
    box(s, x, 2.4, 2.85, 3.4, fill=PANEL, line=CARD, line_w=1.0, radius=True)
    box(s, x + 0.25, 2.7, 0.6, 0.12, fill=col)
    text(s, x + 0.25, 2.95, 2.4, 0.8, [[(title, 17, TEXT, True)]])
    text(s, x + 0.25, 3.7, 2.4, 2.0, [[(body, 13.5, MUTED, False)]])
    x += 3.0

# ==========================================================================
# Slide 4 — Live dashboard (screenshot)
# ==========================================================================
s = slide()
kicker(s, "THE DASHBOARD")
text(s, 0.9, 1.1, 11.5, 0.9, [[("One screen: wall, events, evidence", 30, TEXT, True)]])
shot = os.path.join(SHOTS, "dashboard-full.png")
if os.path.exists(shot):
    box(s, 0.85, 2.05, 11.65, 4.9, fill=PANEL, line=CARD, line_w=1.0)
    s.shapes.add_picture(shot, Inches(0.95), Inches(2.15), width=Inches(11.45))
else:
    box(s, 0.9, 2.1, 11.5, 4.7, fill=PANEL, line=ACCENT, line_w=1.5)
    text(s, 0.9, 4.2, 11.5, 0.8,
         [[("[ add docs/screenshots/dashboard-full.png ]", 18, MUTED, False)]],
         align=PP_ALIGN.CENTER)

# ==========================================================================
# Slide 5 — How it works (pipeline)
# ==========================================================================
s = slide()
kicker(s, "HOW IT WORKS")
text(s, 0.9, 1.1, 11.5, 0.9, [[("A streaming pipeline, GPU-ready", 30, TEXT, True)]])
steps = [
    ("Feeds", "Video files / cameras, looped per source", ACCENT),
    ("Frame batcher", "Short clips sampled for temporal context", ACCENT),
    ("Omni model", "Local VLM scores the watchlist (DGX Spark GPU)", GREEN),
    ("Events + UI", "Live list, snapshots, WebSocket push", AMBER),
]
x = 0.9
for i, (t, b, col) in enumerate(steps):
    box(s, x, 2.7, 2.6, 1.9, fill=PANEL, line=col, line_w=1.25, radius=True)
    text(s, x + 0.2, 2.9, 2.25, 0.6, [[(t, 16, TEXT, True)]])
    text(s, x + 0.2, 3.5, 2.25, 1.0, [[(b, 12.5, MUTED, False)]])
    if i < len(steps) - 1:
        text(s, x + 2.62, 3.25, 0.4, 0.6, [[("→", 26, MUTED, True)]])
    x += 3.0
text(s, 0.9, 5.1, 11.5, 1.3, [
    [("Key design: ", 15, ACCENT, True),
     ("video playback never blocks on the model. Analysis runs on worker "
      "threads; stale work is dropped under load, so the wall stays smooth "
      "while the GPU does the heavy lifting.", 15, TEXT, False)],
], space_after=4)

# ==========================================================================
# Slide 6 — Why DGX Spark
# ==========================================================================
s = slide()
kicker(s, "WHY DGX SPARK")
text(s, 0.9, 1.1, 11.5, 0.9, [[("Edge AI, on-prem, private", 30, TEXT, True)]])
bullets(s, 0.95, 2.3, 11.4, 4, [
    "Runs the vision model locally — footage never leaves the premises.",
    "GPU throughput handles multiple feeds with low latency, no cloud bills.",
    "Swappable model adapter: plug in your 'omni' model via an OpenAI-compatible endpoint.",
    "Config-driven: define feeds and the behavior watchlist in a single YAML file.",
], size=19, gap=16)

# ==========================================================================
# Slide 7 — Tech + status
# ==========================================================================
s = slide()
kicker(s, "TECH & STATUS")
text(s, 0.9, 1.1, 11.5, 0.9, [[("Built and working", 30, TEXT, True)]])
# left: stack
box(s, 0.9, 2.3, 5.6, 4.3, fill=PANEL, line=CARD, line_w=1.0, radius=True)
text(s, 1.2, 2.5, 5.0, 0.5, [[("Stack", 17, ACCENT, True)]])
bullets(s, 1.2, 3.05, 5.0, 3.4, [
    "FastAPI backend (Python)",
    "OpenCV feed looping + MJPEG / frame polling",
    "WebSocket live event push, in-memory store",
    "Single static HTML dashboard — no build step",
    "Swappable omni-model adapter (mock + OpenAI-compatible)",
], size=14.5, gap=11)
# right: status
box(s, 6.8, 2.3, 5.6, 4.3, fill=PANEL, line=CARD, line_w=1.0, radius=True)
text(s, 7.1, 2.5, 5.0, 0.5, [[("Status", 17, GREEN, True)]])
bullets(s, 7.1, 3.05, 5.0, 3.4, [
    "6 live feeds running on the wall",
    "Events + snapshots flowing end-to-end",
    "33 automated tests passing",
    "DGX Spark run guide in the README",
    "Next: wire in the production omni model",
], size=14.5, gap=11)

# ==========================================================================
# Slide 8 — Closing
# ==========================================================================
s = slide()
box(s, 0, 7.32, 13.333, 0.18, fill=GREEN)
text(s, 0.9, 2.6, 11.5, 1.6, [
    [("See the crucial moments. Instantly.", 36, TEXT, True)],
    [("Private, on-prem AI video monitoring on the DGX Spark.", 19, MUTED, False)],
])
text(s, 0.9, 5.6, 11.5, 0.6,
     [[("github.com/wanadzhar913/DGX-Spark-Hackathon-ASB-2026", 14, ACCENT, True)]])

prs.save(OUT)
print(f"Saved: {OUT}")
print(f"Slides: {len(prs.slides.__iter__.__self__._sldIdLst)}")
